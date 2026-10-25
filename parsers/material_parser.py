from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


def _validate_file(file_path: str | Path) -> Path:
    """检查文件是否存在且为普通文件。"""
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"文件不存在：{path}")

    if not path.is_file():
        raise ValueError(f"路径不是文件：{path}")

    return path


def _read_text_file(path: Path) -> str:
    """读取文本文件，依次尝试常见编码。"""
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue

    raise ValueError(f"无法识别文本文件编码：{path}")


def extract_txt_text(file_path: str | Path) -> str:
    """提取 TXT 文件中的文本。"""
    path = _validate_file(file_path)

    if path.suffix.lower() != ".txt":
        raise ValueError("extract_txt_text 只支持 TXT 文件。")

    return _read_text_file(path)


def extract_markdown_text(file_path: str | Path) -> str:
    """提取 Markdown 文件，保留原有标记。"""
    path = _validate_file(file_path)

    if path.suffix.lower() not in {".md", ".markdown"}:
        raise ValueError("extract_markdown_text 只支持 MD、MARKDOWN 文件。")

    return _read_text_file(path)


def extract_pdf_text(file_path: str | Path) -> str:
    """按页提取 PDF 中可读取的文本。"""
    path = _validate_file(file_path)

    if path.suffix.lower() != ".pdf":
        raise ValueError("extract_pdf_text 只支持 PDF 文件。")

    reader = PdfReader(str(path))

    if reader.is_encrypted:
        raise ValueError("PDF 文件已加密，暂时无法读取。")

    page_texts = []

    for page in reader.pages:
        text = page.extract_text()
        if text:
            page_texts.append(text)

    return "\n\n".join(page_texts)


def extract_docx_text(file_path: str | Path) -> str:
    """提取 DOCX 文档中的段落和表格。"""
    path = _validate_file(file_path)

    if path.suffix.lower() != ".docx":
        raise ValueError("extract_docx_text 只支持 DOCX 文件。")

    document = Document(str(path))
    content = []

    # 提取正文段落
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            content.append(text)

    # 提取表格内容
    for table_index, table in enumerate(document.tables, start=1):
        content.append(f"[表格 {table_index}]")

        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            content.append("\t".join(cells))

    return "\n".join(content)


def extract_pptx_text(file_path: str | Path) -> str:
    """按幻灯片提取 PPTX 中的文本框和表格。"""
    path = _validate_file(file_path)

    if path.suffix.lower() != ".pptx":
        raise ValueError("extract_pptx_text 只支持 PPTX 文件。")

    presentation = Presentation(str(path))
    content = []

    for slide_index, slide in enumerate(presentation.slides, start=1):
        content.append(f"[幻灯片 {slide_index}]")

        for shape in slide.shapes:
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    content.append("\t".join(cells))

            elif shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    content.append(text)

    return "\n".join(content)


def extract_xlsx_text(file_path: str | Path) -> str:
    """按工作表和行提取 XLSX 中的单元格内容。"""
    path = _validate_file(file_path)

    if path.suffix.lower() != ".xlsx":
        raise ValueError("extract_xlsx_text 只支持 XLSX 文件。")

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    content = []

    try:
        for worksheet in workbook.worksheets:
            content.append(f"[工作表：{worksheet.title}]")

            for row in worksheet.iter_rows(values_only=True):
                # 将非空单元格转换成文本，保留列之间的位置关系
                cells = [
                    "" if value is None else str(value)
                    for value in row
                ]

                # 跳过整行为空的情况
                if any(cell.strip() for cell in cells):
                    content.append("\t".join(cells).rstrip())

    finally:
        workbook.close()

    return "\n".join(content)


def extract_material_text(file_path: str | Path) -> str:
    """根据文件扩展名，统一提取学习资料中的文本。"""
    path = _validate_file(file_path)
    suffix = path.suffix.lower()

    if suffix == ".txt":
        return extract_txt_text(path)

    if suffix == ".pdf":
        return extract_pdf_text(path)

    if suffix == ".docx":
        return extract_docx_text(path)

    if suffix == ".pptx":
        return extract_pptx_text(path)

    if suffix in {".md", ".markdown"}:
        return extract_markdown_text(path)

    if suffix == ".xlsx":
        return extract_xlsx_text(path)

    raise ValueError(
        "当前支持的格式：TXT、PDF、DOCX、PPTX、MD、MARKDOWN、XLSX。"
    )
