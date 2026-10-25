
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from storage.db import init_db
from services.material_service import (
    import_material,
    list_materials,
    get_material,
    delete_material,
)


class TestMaterialImportFormats(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)

        self.db_path = self.root / "test.db"
        self.materials_dir = self.root / "materials"
        self.source_dir = self.root / "source"
        self.source_dir.mkdir()

        init_db(db_path=self.db_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def _create_source_file(self, suffix):
        """为不同格式创建真实的测试文件。"""
        path = self.source_dir / f"sample{suffix}"

        if suffix == ".txt":
            path.write_text("TXT 测试内容", encoding="utf-8")

        elif suffix == ".md":
            path.write_text(
                "# 复习提纲\n\n- 进程管理",
                encoding="utf-8",
            )

        elif suffix == ".docx":
            document = Document()
            document.add_paragraph("DOCX 测试内容")
            document.save(path)

        elif suffix == ".pptx":
            presentation = Presentation()
            slide = presentation.slides.add_slide(
                presentation.slide_layouts[5]
            )
            textbox = slide.shapes.add_textbox(
                Inches(1), Inches(1), Inches(5), Inches(1)
            )
            textbox.text = "PPTX 测试内容"
            presentation.save(path)

        elif suffix == ".xlsx":
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "测试表"
            worksheet.append(["知识点", "状态"])
            worksheet.append(["数据结构", "待复习"])
            workbook.save(path)
            workbook.close()

        elif suffix == ".pdf":
            # 仅测试材料服务的导入流程，PDF 文本提取由独立测试覆盖。
            path.write_bytes(b"%PDF-test-placeholder")

        else:
            raise ValueError(f"未配置测试文件格式：{suffix}")

        return path

    def _assert_import_query_delete(self, suffix, expected_text):
        source_path = self._create_source_file(suffix)

        if suffix == ".pdf":
            # 只模拟 PDF 解析结果，不模拟文件复制或数据库操作。
            with patch(
                "services.material_service.extract_material_text",
                return_value=expected_text,
            ):
                import_material(
                    source_path,
                    db_path=self.db_path,
                    materials_dir=self.materials_dir,
                )
        else:
            import_material(
                source_path,
                db_path=self.db_path,
                materials_dir=self.materials_dir,
            )

        # 1. 确认数据库中有且只有一条记录
        records = list_materials(db_path=self.db_path)
        self.assertEqual(len(records), 1)

        material_id = records[0]["id"]

        # 2. 确认详情、文本及存储文件
        material = get_material(material_id, db_path=self.db_path)
        self.assertIsNotNone(material)
        self.assertEqual(material["file_type"], suffix.lstrip("."))
        self.assertTrue(material["extracted_text"].strip())

        if suffix == ".pdf":
            self.assertEqual(material["extracted_text"], expected_text)
        else:
            self.assertIn(expected_text, material["extracted_text"])

        stored_path = Path(material["stored_path"])
        self.assertTrue(stored_path.is_file())
        self.assertEqual(stored_path.suffix.lower(), suffix)

        # 3. 确认删除后，记录与存储文件都不存在
        self.assertTrue(
            delete_material(
                material_id,
                db_path=self.db_path,
                materials_dir=self.materials_dir,
            )
        )
        self.assertFalse(stored_path.exists())
        self.assertEqual(
            list_materials(db_path=self.db_path),
            [],
        )

    def test_import_txt(self):
        self._assert_import_query_delete(".txt", "TXT 测试内容")

    def test_import_markdown(self):
        self._assert_import_query_delete(".md", "进程管理")

    def test_import_docx(self):
        self._assert_import_query_delete(".docx", "DOCX 测试内容")

    def test_import_pptx(self):
        self._assert_import_query_delete(".pptx", "PPTX 测试内容")

    def test_import_xlsx(self):
        self._assert_import_query_delete(".xlsx", "数据结构")

    def test_import_pdf(self):
        self._assert_import_query_delete(".pdf", "PDF 模拟提取内容")


if __name__ == "__main__":
    unittest.main()
