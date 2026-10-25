
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from parsers.material_parser import extract_material_text


class TestOfficeParsers(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_docx_parser(self):
        path = self.root / "notes.docx"

        document = Document()
        document.add_paragraph("操作系统复习提纲")

        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "知识点"
        table.cell(0, 1).text = "状态"
        table.cell(1, 0).text = "进程调度"
        table.cell(1, 1).text = "待复习"
        document.save(path)

        text = extract_material_text(path)

        self.assertIn("操作系统复习提纲", text)
        self.assertIn("知识点", text)
        self.assertIn("进程调度", text)
        self.assertIn("待复习", text)

    def test_pptx_parser(self):
        path = self.root / "lesson.pptx"

        presentation = Presentation()
        slide = presentation.slides.add_slide(
            presentation.slide_layouts[5]
        )

        textbox = slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(5), Inches(1)
        )
        textbox.text = "第三章：进程管理"

        table_shape = slide.shapes.add_table(
            2, 2, Inches(1), Inches(2), Inches(5), Inches(1)
        )
        table_shape.table.cell(0, 0).text = "概念"
        table_shape.table.cell(0, 1).text = "说明"
        table_shape.table.cell(1, 0).text = "进程"
        table_shape.table.cell(1, 1).text = "资源分配单位"

        presentation.save(path)

        text = extract_material_text(path)

        self.assertIn("第三章：进程管理", text)
        self.assertIn("进程", text)
        self.assertIn("资源分配单位", text)
        self.assertIn("[幻灯片 1]", text)

    def test_markdown_parser(self):
        path = self.root / "notes.md"
        original = "# 数据结构复习\n\n- 栈\n- 队列\n"

        path.write_text(original, encoding="utf-8")

        text = extract_material_text(path)

        self.assertEqual(text, original)

    def test_xlsx_parser(self):
        path = self.root / "outline.xlsx"

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "复习提纲"
        worksheet.append(["知识点", "掌握情况"])
        worksheet.append(["二叉树", "已掌握"])
        worksheet.append(["图的遍历", "待复习"])
        workbook.save(path)
        workbook.close()

        text = extract_material_text(path)

        self.assertIn("[工作表：复习提纲]", text)
        self.assertIn("知识点", text)
        self.assertIn("二叉树", text)
        self.assertIn("已掌握", text)
        self.assertIn("图的遍历", text)
        self.assertIn("待复习", text)

    def test_uppercase_extension(self):
        # 扩展名大小写不应影响解析
        path = self.root / "notes.MD"
        path.write_text("# 复习资料", encoding="utf-8")

        text = extract_material_text(path)

        self.assertIn("# 复习资料", text)


if __name__ == "__main__":
    unittest.main()
